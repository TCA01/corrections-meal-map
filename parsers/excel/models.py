from __future__ import annotations

from dataclasses import asdict, dataclass, field
from typing import Any


@dataclass
class CellData:
    coordinate: str
    row: int
    column: int
    raw_value: Any
    text: str
    number_format: str | None = None
    is_formula: bool = False
    hidden_row: bool = False
    hidden_column: bool = False
    merged_parent: str | None = None
    merged_range: str | None = None
    formula: str | None = None
    display_value: Any = None
    formula_cache_status: str | None = None
    is_excel_error: bool = False

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass
class SheetGrid:
    name: str
    actual_min_row: int | None
    actual_max_row: int | None
    actual_min_col: int | None
    actual_max_col: int | None
    merged_ranges: list[dict[str, Any]] = field(default_factory=list)
    cells: list[CellData] = field(default_factory=list)

    def by_position(self) -> dict[tuple[int, int], CellData]:
        return {(cell.row, cell.column): cell for cell in self.cells}

    def cell(self, row: int, column: int) -> CellData | None:
        return self.by_position().get((row, column))

    def effective_text(self, row: int, column: int) -> str:
        cell = self.cell(row, column)
        if cell is None:
            return ""
        if cell.text:
            return cell.text
        if cell.merged_parent:
            parent = next((item for item in self.cells if item.coordinate == cell.merged_parent), None)
            return parent.text if parent else ""
        return ""

    def to_dict(self) -> dict[str, Any]:
        return {
            "sheet": self.name,
            "actual_range": self.actual_range,
            "actual_min_row": self.actual_min_row,
            "actual_max_row": self.actual_max_row,
            "actual_min_col": self.actual_min_col,
            "actual_max_col": self.actual_max_col,
            "merged_ranges": self.merged_ranges,
            "cells": [cell.to_dict() for cell in self.cells],
        }

    @property
    def actual_range(self) -> str | None:
        if self.actual_min_row is None:
            return None
        from openpyxl.utils import get_column_letter

        return (
            f"{get_column_letter(self.actual_min_col)}{self.actual_min_row}:"
            f"{get_column_letter(self.actual_max_col)}{self.actual_max_row}"
        )


@dataclass
class WorkbookGrid:
    path: str
    declared_extension: str
    detected_format: str
    format_extension_mismatch: bool
    sheets: list[SheetGrid]

    def to_dict(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "declared_extension": self.declared_extension,
            "detected_format": self.detected_format,
            "format_extension_mismatch": self.format_extension_mismatch,
            "sheets": [sheet.to_dict() for sheet in self.sheets],
        }


@dataclass
class MenuItem:
    name: str
    raw_text: str

    def to_dict(self) -> dict[str, str]:
        return asdict(self)


@dataclass
class MealRecord:
    institution_id: str
    meal_date: str
    meal_type: str
    menu_items: list[MenuItem]
    source: dict[str, Any]
    parser: dict[str, Any]
    schema_version: str = "1.0"
    validation_status: str = "valid"
    validation_issues: list[dict[str, str]] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        result = asdict(self)
        result["menu_items"] = [item.to_dict() for item in self.menu_items]
        return result


@dataclass
class ParseFailure:
    code: str
    message: str
    sheet: str | None = None

    def to_dict(self) -> dict[str, str | None]:
        return asdict(self)

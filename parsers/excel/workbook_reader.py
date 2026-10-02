from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
import math

import openpyxl
import xlrd
from openpyxl.utils import get_column_letter

from .detector import detect_excel_format
from .models import CellData, SheetGrid, WorkbookGrid
from .menu_artifacts import EXCEL_ERRORS


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def read_workbook(path: Path, declared_extension: str | None = None) -> WorkbookGrid:
    detected = detect_excel_format(path)
    declared = (declared_extension or path.suffix.lstrip(".")).lower()
    sheets = _read_xlsx(path) if detected == "xlsx" else _read_xls(path)
    return WorkbookGrid(
        path=path.as_posix(),
        declared_extension=declared,
        detected_format=detected,
        format_extension_mismatch=declared != detected,
        sheets=sheets,
    )


def _read_xlsx(path: Path) -> list[SheetGrid]:
    content = path.read_bytes()
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False, read_only=False)
    formula_positions = {worksheet.title: {cell.coordinate for row in worksheet.iter_rows() for cell in row
                                           if cell.data_type == "f"}
                         for worksheet in workbook.worksheets}
    cached_values = {}
    cached_errors = set()
    if any(formula_positions.values()):
        cached = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
        try:
            for worksheet in cached.worksheets:
                if formula_positions[worksheet.title]:
                    cached_values[worksheet.title] = {}
                    for cached_row in worksheet.iter_rows():
                        for cached_cell in cached_row:
                            coordinate = getattr(cached_cell, "coordinate", None)
                            if coordinate in formula_positions[worksheet.title]:
                                cached_values[worksheet.title][coordinate] = cached_cell.value
                                if cached_cell.data_type == "e":
                                    cached_errors.add((worksheet.title, coordinate))
        finally:
            cached.close()
    results: list[SheetGrid] = []
    for worksheet in workbook.worksheets:
        merged_lookup: dict[tuple[int, int], tuple[str, str]] = {}
        meaningful_merges = []
        for merged in worksheet.merged_cells.ranges:
            master = worksheet.cell(merged.min_row, merged.min_col)
            if master.value in (None, ""):
                continue
            covered = []
            for row in range(merged.min_row, merged.max_row + 1):
                for column in range(merged.min_col, merged.max_col + 1):
                    coordinate = f"{get_column_letter(column)}{row}"
                    covered.append(coordinate)
                    merged_lookup[(row, column)] = (master.coordinate, str(merged))
            meaningful_merges.append({
                "range": str(merged),
                "master_coordinate": master.coordinate,
                "covered_coordinates": covered,
                "merged_value": _text(master.value),
            })
        positions = {
            (cell.row, cell.column)
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value not in (None, "")
        }
        positions.update(merged_lookup)
        cells = []
        for row, column in sorted(positions):
            cell = worksheet.cell(row, column)
            formula = cell.value if cell.data_type == "f" else None
            cached_value = cached_values.get(worksheet.title, {}).get(cell.coordinate) if formula else None
            cache_valid = (formula is not None and isinstance(cached_value, (str, int, float, bool, date, datetime))
                           and _text(cached_value) != "" and _text(cached_value).upper() not in EXCEL_ERRORS
                           and not _text(cached_value).startswith("=")
                           and (worksheet.title, cell.coordinate) not in cached_errors
                           and (not isinstance(cached_value, float) or math.isfinite(cached_value)))
            display = cached_value if cache_valid else cell.value
            if formula and (_text(cached_value).upper() in EXCEL_ERRORS or (worksheet.title, cell.coordinate) in cached_errors):
                display = cached_value
            parent, merged_range = merged_lookup.get((row, column), (None, None))
            cells.append(CellData(
                coordinate=cell.coordinate,
                row=row,
                column=column,
                raw_value=cell.value,
                text=_text(display),
                number_format=cell.number_format,
                is_formula=formula is not None,
                formula=formula,
                display_value=display,
                formula_cache_status=("resolved" if cache_valid else "unresolved") if formula else None,
                is_excel_error=cell.data_type == "e" or (worksheet.title, cell.coordinate) in cached_errors,
                hidden_row=bool(worksheet.row_dimensions[row].hidden),
                hidden_column=bool(worksheet.column_dimensions[get_column_letter(column)].hidden),
                merged_parent=parent if parent != cell.coordinate else None,
                merged_range=merged_range,
            ))
        results.append(_sheet(worksheet.title, cells, meaningful_merges))
    workbook.close()
    return results


def _read_xls(path: Path) -> list[SheetGrid]:
    workbook = xlrd.open_workbook(path, formatting_info=True)
    results = []
    for worksheet in workbook.sheets():
        merged_lookup: dict[tuple[int, int], tuple[str, str]] = {}
        meaningful_merges = []
        for rlo, rhi, clo, chi in worksheet.merged_cells:
            master_value = worksheet.cell_value(rlo, clo)
            if master_value in (None, ""):
                continue
            master = f"{get_column_letter(clo + 1)}{rlo + 1}"
            range_name = f"{master}:{get_column_letter(chi)}{rhi}"
            covered = []
            for row in range(rlo, rhi):
                for column in range(clo, chi):
                    coordinate = f"{get_column_letter(column + 1)}{row + 1}"
                    covered.append(coordinate)
                    merged_lookup[(row + 1, column + 1)] = (master, range_name)
            meaningful_merges.append({
                "range": range_name,
                "master_coordinate": master,
                "covered_coordinates": covered,
                "merged_value": _text(master_value),
            })
        positions = {
            (row + 1, column + 1)
            for row in range(worksheet.nrows)
            for column in range(worksheet.ncols)
            if worksheet.cell_value(row, column) not in (None, "")
        }
        positions.update(merged_lookup)
        cells = []
        for row, column in sorted(positions):
            cell = worksheet.cell(row - 1, column - 1)
            value: Any = cell.value
            if cell.ctype == xlrd.XL_CELL_ERROR:
                value = xlrd.error_text_from_code.get(int(cell.value), "#VALUE!")
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    value = xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode)
                except (ValueError, OverflowError):
                    pass
            parent, merged_range = merged_lookup.get((row, column), (None, None))
            coordinate = f"{get_column_letter(column)}{row}"
            cells.append(CellData(
                coordinate=coordinate,
                row=row,
                column=column,
                raw_value=value,
                text=_text(value),
                number_format=None,
                is_formula=False,
                is_excel_error=cell.ctype == xlrd.XL_CELL_ERROR,
                merged_parent=parent if parent != coordinate else None,
                merged_range=merged_range,
            ))
        results.append(_sheet(worksheet.name, cells, meaningful_merges))
    return results


def _sheet(name: str, cells: list[CellData], merges: list[dict[str, Any]]) -> SheetGrid:
    meaningful = [cell for cell in cells if cell.text or cell.merged_parent]
    if not meaningful:
        return SheetGrid(name, None, None, None, None, merges, cells)
    return SheetGrid(
        name=name,
        actual_min_row=min(cell.row for cell in meaningful),
        actual_max_row=max(cell.row for cell in meaningful),
        actual_min_col=min(cell.column for cell in meaningful),
        actual_max_col=max(cell.column for cell in meaningful),
        merged_ranges=merges,
        cells=cells,
    )

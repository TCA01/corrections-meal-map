from __future__ import annotations

import re
import zipfile
from collections import Counter
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree


MEAL_KEYWORDS = ("아침", "조식", "점심", "중식", "저녁", "석식")
DATE_PATTERN = re.compile(r"(?:20\d{2}[년./-])?\s*(?:1[0-2]|0?[1-9])\s*월|\b\d{1,2}[./-]\d{1,2}\b")


def survey_file(path: Path) -> dict[str, object]:
    extension = path.suffix.lower().lstrip(".")
    base: dict[str, object] = {
        "path": path.as_posix(),
        "filename": path.name,
        "extension": extension,
        "file_size": path.stat().st_size,
    }
    try:
        if extension == "xlsx":
            base.update(survey_xlsx(path))
        elif extension == "xls":
            base.update(survey_xls(path))
        elif extension == "hwpx":
            base.update(survey_hwpx(path))
        elif extension == "pdf":
            base.update(survey_pdf(path))
        elif extension == "hwp":
            base.update(survey_hwp(path))
        else:
            base.update({"survey_status": "unsupported", "reason": f"unsupported extension: {extension}"})
    except Exception as exc:
        base.update({"survey_status": "failed", "error_type": type(exc).__name__, "error": str(exc)})
    return base


def survey_xlsx(path: Path) -> dict[str, object]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(path.read_bytes()), read_only=False, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        date_like = 0
        keyword_counts = {keyword: 0 for keyword in MEAL_KEYWORDS}
        nonempty = 0
        min_row = min_column = None
        max_used_row = max_used_column = 0
        formula_cells = 0
        numeric_cells = 0
        formatted_date_cells = 0
        number_formats: Counter[str] = Counter()
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                nonempty += 1
                min_row = cell.row if min_row is None else min(min_row, cell.row)
                min_column = cell.column if min_column is None else min(min_column, cell.column)
                max_used_row = max(max_used_row, cell.row)
                max_used_column = max(max_used_column, cell.column)
                text = str(cell.value)
                formula_cells += cell.data_type == "f"
                numeric_cells += isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
                formatted_date_cells += bool(cell.is_date)
                if cell.number_format and cell.number_format != "General":
                    number_formats[cell.number_format] += 1
                if DATE_PATTERN.search(text):
                    date_like += 1
                for keyword in MEAL_KEYWORDS:
                    if keyword in text:
                        keyword_counts[keyword] += 1
        sheets.append({
            "name": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "nonempty_cells": nonempty,
            "nonempty_used_range": {
                "min_row": min_row,
                "min_column": min_column,
                "max_row": max_used_row or None,
                "max_column": max_used_column or None,
            },
            "merged_cell_ranges": len(sheet.merged_cells.ranges),
            "hidden_rows": sum(1 for dimension in sheet.row_dimensions.values() if dimension.hidden),
            "hidden_columns": sum(1 for dimension in sheet.column_dimensions.values() if dimension.hidden),
            "date_like_cells": date_like,
            "formula_cells": formula_cells,
            "numeric_cells": numeric_cells,
            "formatted_date_cells": formatted_date_cells,
            "number_formats": dict(number_formats.most_common(12)),
            "meal_keyword_counts": keyword_counts,
        })
    workbook.close()
    return {"survey_status": "ok", "valid_workbook": True, "sheet_count": len(sheets), "sheets": sheets}


def survey_xls(path: Path) -> dict[str, object]:
    if path.read_bytes()[:4] == b"PK\x03\x04":
        result = survey_xlsx(path)
        result.update({
            "declared_format": "xls",
            "detected_format": "xlsx",
            "format_extension_mismatch": True,
        })
        return result
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    sheets = []
    for sheet in workbook.sheets():
        keyword_counts = {keyword: 0 for keyword in MEAL_KEYWORDS}
        date_like = numeric_cells = nonempty = 0
        for row in range(sheet.nrows):
            for column in range(sheet.ncols):
                cell = sheet.cell(row, column)
                if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                    continue
                nonempty += 1
                numeric_cells += cell.ctype == xlrd.XL_CELL_NUMBER
                date_like += cell.ctype == xlrd.XL_CELL_DATE or bool(DATE_PATTERN.search(str(cell.value)))
                text = str(cell.value)
                for keyword in MEAL_KEYWORDS:
                    keyword_counts[keyword] += text.count(keyword)
        sheets.append({
            "name": sheet.name,
            "max_row": sheet.nrows,
            "max_column": sheet.ncols,
            "nonempty_cells": nonempty,
            "nonempty_used_range": {
                "min_row": 1 if sheet.nrows else None,
                "min_column": 1 if sheet.ncols else None,
                "max_row": sheet.nrows or None,
                "max_column": sheet.ncols or None,
            },
            "merged_cell_ranges": len(sheet.merged_cells),
            "hidden_rows": 0,
            "hidden_columns": 0,
            "formula_cells": None,
            "numeric_cells": numeric_cells,
            "formatted_date_cells": date_like,
            "date_like_cells": date_like,
            "number_formats": {},
            "meal_keyword_counts": keyword_counts,
        })
    workbook.release_resources()
    return {
        "survey_status": "ok",
        "valid_workbook": True,
        "declared_format": "xls",
        "detected_format": "xls",
        "format_extension_mismatch": False,
        "sheet_count": len(sheets),
        "sheets": sheets,
    }


def survey_hwpx(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        section_names = sorted(
            name for name in archive.namelist() if re.search(r"(?:^|/)section\d+\.xml$", name, re.I)
        )
        table_count = 0
        nested_table_count = 0
        row_count = 0
        cell_count = 0
        max_columns = 0
        text_parts: list[str] = []
        for name in section_names:
            root = ElementTree.fromstring(archive.read(name))
            def walk(node: ElementTree.Element, table_depth: int = 0) -> None:
                nonlocal table_count, nested_table_count, row_count, cell_count, max_columns
                local = node.tag.rsplit("}", 1)[-1].lower()
                next_depth = table_depth
                if local in {"tbl", "table"}:
                    table_count += 1
                    nested_table_count += table_depth > 0
                    next_depth += 1
                if local in {"tr", "row"}:
                    row_count += 1
                    direct_cells = sum(
                        child.tag.rsplit("}", 1)[-1].lower() in {"tc", "cell"} for child in list(node)
                    )
                    max_columns = max(max_columns, direct_cells)
                if local in {"tc", "cell"}:
                    cell_count += 1
                for child in list(node):
                    walk(child, next_depth)
            walk(root)
            for node in root.iter():
                if node.text:
                    text_parts.append(node.text)
        text = " ".join(text_parts)
    return {
        "survey_status": "ok",
        "valid_zip": bad_member is None,
        "section_count": len(section_names),
        "table_count": table_count,
        "nested_table_count": nested_table_count,
        "row_count": row_count,
        "cell_count": cell_count,
        "max_detected_columns": max_columns,
        "text_character_count": len(text),
        "meal_keyword_counts": {keyword: text.count(keyword) for keyword in MEAL_KEYWORDS},
        "date_keyword_count": len(DATE_PATTERN.findall(text)),
    }


def survey_pdf(path: Path) -> dict[str, object]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    page_reports = []
    total_text = 0
    total_images = 0
    table_like_lines = 0
    for index, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        text_blocks = 0
        def visitor(fragment: str, cm: object, tm: object, font: object, size: object) -> None:
            nonlocal text_blocks
            if fragment.strip():
                text_blocks += 1
        try:
            page.extract_text(visitor_text=visitor)
        except TypeError:
            text_blocks = len([line for line in text.splitlines() if line.strip()])
        image_count = len(getattr(page, "images", []))
        total_text += len(text.strip())
        total_images += image_count
        line_count = sum(1 for line in text.splitlines() if re.search(r"\S+\s{2,}\S+", line))
        table_like_lines += line_count
        page_reports.append({
            "page": index,
            "text_characters": len(text.strip()),
            "image_count": image_count,
            "table_like_lines": line_count,
            "width_points": round(float(page.mediabox.width), 2),
            "height_points": round(float(page.mediabox.height), 2),
            "text_block_count": text_blocks,
        })
    return {
        "survey_status": "ok",
        "page_count": len(reader.pages),
        "text_character_count": total_text,
        "has_text_layer": total_text > 20,
        "image_count": total_images,
        "possible_image_only": total_text <= 20 and total_images > 0,
        "table_like_line_count": table_like_lines,
        "pages": page_reports,
    }


def survey_hwp(path: Path) -> dict[str, object]:
    signature = path.read_bytes()[:8]
    ole_signature = bytes.fromhex("D0CF11E0A1B11AE1")
    is_ole = signature == ole_signature
    streams: list[str] = []
    parser_available = False
    parser_error = None
    try:
        import olefile

        parser_available = True
        if is_ole:
            with olefile.OleFileIO(path) as ole:
                streams = ["/".join(parts) for parts in ole.listdir()]
    except Exception as exc:
        parser_error = str(exc)
    return {
        "survey_status": "ok" if is_ole else "failed",
        "signature_hex": signature.hex(),
        "is_ole_cfb": is_ole,
        "local_ole_parser_available": parser_available,
        "ole_stream_count": len(streams),
        "ole_streams": streams[:40],
        "parser_error": parser_error,
    }


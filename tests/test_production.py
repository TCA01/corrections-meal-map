from __future__ import annotations

import hashlib
import json
from dataclasses import replace
from pathlib import Path

import pytest

from collector.models import Attachment, Post
from config.settings import Settings
from production.downloader import ProductionExcelDownloader
from production.gate import production_route
from production.io import atomic_json, atomic_jsonl, read_jsonl
from production.manifest import build_excel_manifest
from production.pipeline import ProductionExcelPipeline, detect_conflicts
from production.validation import validate_dataset
from parsers.excel.models import WorkbookGrid, SheetGrid, CellData
from parsers.excel.pipeline import ExcelMealParser


FIXTURES = Path(__file__).parent / "fixtures"


def settings(tmp_path: Path) -> Settings:
    return replace(
        Settings(), project_root=tmp_path,
        catalog_path=tmp_path / "catalog/posts.jsonl",
        historical_catalog_path=tmp_path / "audit/historical.jsonl",
        parser_samples_path=tmp_path / "audit/samples.json",
        institutions_path=tmp_path / "institutions.json",
        production_root=tmp_path / "production",
        production_manifest_path=tmp_path / "production/manifest.jsonl",
        production_download_checkpoint_path=tmp_path / "production/checkpoints/download.json",
        production_parse_checkpoint_path=tmp_path / "production/checkpoints/parse.json",
        production_raw_root=tmp_path / "production/raw",
    )


def manifest_item(document_id: str = "1-1", *, institution_id: str | None = "I1") -> dict:
    post_id, attachment_id = document_id.split("-")
    return {
        "document_id": document_id, "post_id": post_id, "attachment_id": attachment_id,
        "institution_id": institution_id, "institution_name": "기관", "filename": f"{document_id}.xlsx",
        "declared_extension": "xlsx", "extension": "xlsx", "document_role": "inmate",
        "meal_year": 2026, "meal_month": 1, "published_date": "2025-12-20",
        "post_url": f"https://example.test/post/{post_id}",
        "download_url": f"https://example.test/download/{attachment_id}",
        "download_status": "reused", "sha256": hashlib.sha256(b"x").hexdigest(), "file_size": 1,
        "local_path": f"raw/{document_id}.xlsx", "duplicate_of": None,
        "parser_status": "pending", "production_status": "pending", "error": None,
    }


def ready_result(document_id: str = "1-1", *, menu: str = "밥", institution_id: str = "I1") -> dict:
    post_id, attachment_id = document_id.split("-")
    return {
        "sample_id": document_id, "post_id": post_id, "attachment_id": attachment_id,
        "filename": f"{document_id}.xlsx", "institution_id": institution_id,
        "declared_extension": "xlsx", "detected_format": "xlsx", "format_extension_mismatch": False,
        "meal_year": 2026, "meal_month": 1, "period_evidence": ["metadata"],
        "layout_profiles": [{"detected_layout_family": "date_rows_meal_columns"}],
        "status": "PASS", "records_generated": 1, "valid": 1, "warnings": 0, "invalid": 0,
        "unresolved": 0, "coverage": {"complete": True, "missing_dates": [], "missing_meals": []},
        "source_provenance_valid": True, "document_issues": [], "failures": [],
        "menu_quality_valid": True, "unresolved_formula_values": 0,
        "excel_error_menu_items": 0, "unresolved_non_menu_artifacts": 0,
        "records": [{
            "institution_id": institution_id, "meal_date": "2026-01-01", "meal_type": "breakfast",
            "menu_items": [{"name": menu, "raw_text": menu}],
            "source": {"post_id": post_id, "attachment_id": attachment_id, "sheet": "Sheet1", "cells": ["B2"]},
            "parser": {"format": "xlsx", "layout_family": "date_rows_meal_columns", "parser_version": "3A.1"},
            "schema_version": "1.0", "validation_status": "valid", "validation_issues": [],
        }],
    }


def write_institution(path: Path) -> None:
    atomic_json(path, [{
        "institution_id": "I1", "canonical_name": "기관", "short_name": "기관",
        "institution_type": "prison", "aliases": [], "address": "", "source_url": "https://example.test",
        "active": True, "postal_code": None, "phone": None, "latitude": None, "longitude": None,
    }])


def write_inputs(cfg: Settings, items: list[dict]) -> None:
    write_institution(cfg.institutions_path)
    for item in items:
        path = cfg.project_root / item["local_path"]
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"x")
    atomic_jsonl(cfg.production_manifest_path, items)


@pytest.mark.parametrize("case", json.loads((FIXTURES / "production_gate/cases.json").read_text(encoding="utf-8")), ids=lambda c: c["name"])
def test_production_gate_fixture(case: dict) -> None:
    result = {
        "status": case["status"], "invalid": case["invalid"], "failures": [],
        "coverage": {"complete": True, "missing_dates": case["missing_dates"], "missing_meals": case["missing_meals"]},
        "source_provenance_valid": case["provenance"], "document_issues": [],
        "menu_quality_valid": True,
    }
    assert production_route(result, institution_id=case["institution_id"]) == case["expected"]


def test_year_inference_is_review() -> None:
    result = ready_result()
    result["document_issues"] = [{"code": "YEAR_INFERRED_FROM_PUBLICATION"}]
    assert production_route(result, institution_id="I1") == "review"


def test_manifest_filters_role_and_excel(tmp_path: Path) -> None:
    cfg = settings(tmp_path)
    post = Post("1", "기관", "식단", "2026-01-01", "https://example.test", institution_id="I1",
                attachment_count=3, is_meal_plan=True, meal_year=2026, meal_month=1, attachments=[
        Attachment("1", "a.xlsx", "xlsx", "https://example.test/1", document_role="inmate"),
        Attachment("2", "b.pdf", "pdf", "https://example.test/2", document_role="inmate"),
        Attachment("3", "c.xls", "xls", "https://example.test/3", document_role="staff"),
    ])
    atomic_jsonl(cfg.historical_catalog_path, [post.to_dict()])
    assert [item["attachment_id"] for item in build_excel_manifest(cfg)] == ["1"]


def test_download_reuses_verified_file(tmp_path: Path) -> None:
    cfg = settings(tmp_path)
    path = tmp_path / "raw/1-1.xlsx"; path.parent.mkdir(); path.write_bytes(b"x")
    item = manifest_item(); item.update(sha256=hashlib.sha256(b"x").hexdigest(), file_size=1)
    atomic_jsonl(cfg.production_manifest_path, [item])
    result = ProductionExcelDownloader(cfg).run()
    assert result.reused == 1 and result.completed


def test_download_checkpoint_from_index(tmp_path: Path) -> None:
    cfg = settings(tmp_path)
    items = [manifest_item("1-1"), manifest_item("2-2")]
    for item in items:
        path = tmp_path / item["local_path"]; path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(b"x")
        item.update(sha256=hashlib.sha256(b"x").hexdigest(), file_size=1)
    atomic_jsonl(cfg.production_manifest_path, items)
    result = ProductionExcelDownloader(cfg).run(from_checkpoint=1, limit=1)
    assert result.next_index == 2 and json.loads(cfg.production_download_checkpoint_path.read_text())["completed"]


def test_parse_checkpoint_resume(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = settings(tmp_path); write_inputs(cfg, [manifest_item("1-1"), manifest_item("2-2")])
    pipeline = ProductionExcelPipeline(cfg); calls = []
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda item: calls.append(item["document_id"]) or ready_result(item["document_id"]))
    result = pipeline.run(limit=1)
    assert result["completed"] is False and calls == ["1-1"]
    checkpoint = json.loads(cfg.production_parse_checkpoint_path.read_text())
    assert checkpoint["next_index"] == 1
    pipeline.run()
    assert calls == ["1-1", "2-2"]


def test_routing_ready_review_failed_and_unresolved() -> None:
    ready = ready_result()
    review = ready_result(); review["status"] = "PARTIAL"
    failed = ready_result(); failed["status"] = "FAIL"; failed["failures"] = [{"code": "CORRUPT_WORKBOOK"}]
    assert production_route(ready, institution_id="I1") == "ready"
    assert production_route(review, institution_id="I1") == "review"
    assert production_route(failed, institution_id="I1") == "failed"
    assert production_route(ready, institution_id=None) == "review"


def test_identical_content_is_deduplicated() -> None:
    docs = [{"document_id": i, "result": ready_result(i)} for i in ("1-1", "2-2")]
    affected, conflicts, duplicates = detect_conflicts(docs)
    assert not affected and not conflicts and duplicates == 1


def test_different_content_is_conflict() -> None:
    docs = [{"document_id": "1-1", "result": ready_result("1-1", menu="밥")},
            {"document_id": "2-2", "result": ready_result("2-2", menu="죽")}]
    affected, conflicts, _ = detect_conflicts(docs)
    assert affected == {"1-1", "2-2"} and conflicts[0]["code"] == "DATA_CONFLICT"


def test_new_layout_is_failed_quarantine() -> None:
    result = ready_result(); result.update(status="FAIL", records=[], failures=[{"code": "UNSUPPORTED_LAYOUT"}])
    assert production_route(result, institution_id="I1") == "failed"


def publish_one(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[Settings, dict, Path]:
    cfg = settings(tmp_path); write_institution(cfg.institutions_path)
    write_inputs(cfg, [manifest_item()])
    pipeline = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda item: ready_result())
    report = pipeline.run(new_run=True)
    current = json.loads((cfg.production_root / "current.json").read_text())
    return cfg, report, cfg.production_root / current["run_path"]


def test_staging_atomic_publish(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg, report, run = publish_one(tmp_path, monkeypatch)
    assert report["atomic_publication"] == "PASS" and run.is_dir()
    assert not (cfg.production_root / ".staging" / report["run_metadata"]["dataset_version"]).exists()


def test_failed_publish_preserves_previous_pointer(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = settings(tmp_path); write_institution(cfg.institutions_path)
    atomic_json(cfg.production_root / "current.json", {"dataset_version": "old", "run_path": "runs/old"})
    write_inputs(cfg, [manifest_item()])
    pipeline = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda item: ready_result())
    monkeypatch.setattr(pipeline, "_validate_staging", lambda *args: (_ for _ in ()).throw(ValueError("bad")))
    with pytest.raises(ValueError): pipeline.run(new_run=True)
    assert json.loads((cfg.production_root / "current.json").read_text())["dataset_version"] == "old"


def test_public_projection_has_urls_without_local_path(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _, _, run = publish_one(tmp_path, monkeypatch)
    text = next((run / "web_data/menus").rglob("*.json")).read_text(encoding="utf-8")
    assert "https://example.test/post/1" in text and "download/1" in text and "local_path" not in text


def test_institution_year_month_manifest(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _, _, run = publish_one(tmp_path, monkeypatch)
    value = json.loads((run / "web_data/manifest.json").read_text(encoding="utf-8"))
    assert value["institutions"][0]["available_years"] == {"2026": [1]}


def test_production_dataset_schema_and_month_shape(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _, report, run = publish_one(tmp_path, monkeypatch)
    month = json.loads(next((run / "web_data/menus").rglob("*.json")).read_text(encoding="utf-8"))
    assert report["run_metadata"]["schema_version"] == "1.0"
    assert month["schema_version"] == "1.0" and month["days"]["2026-01-01"]["breakfast"]["source_document_id"] == "1-1"


def test_report_contains_metrics(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _, report, _ = publish_one(tmp_path, monkeypatch)
    for key in ("documents", "records", "formats", "layout_families", "institution_coverage", "date_coverage", "web_data"):
        assert key in report


def test_missing_coverage_cannot_be_ready() -> None:
    result = ready_result()
    del result["coverage"]
    assert production_route(result, institution_id="I1") == "review"


def test_parse_resume_rejects_changed_inputs(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = settings(tmp_path); write_inputs(cfg, [manifest_item("1-1"), manifest_item("2-2")])
    pipeline = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda i: ready_result(i["document_id"]))
    pipeline.run(limit=1)
    items = read_jsonl(cfg.production_manifest_path); items.reverse()
    atomic_jsonl(cfg.production_manifest_path, items)
    with pytest.raises(ValueError, match="inputs changed"):
        pipeline.run()


def test_completed_parse_is_idempotent(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg, report, _ = publish_one(tmp_path, monkeypatch)
    pipeline = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda i: pytest.fail("must reuse completed run"))
    assert pipeline.run()["run_metadata"] == report["run_metadata"]


def test_conflicting_sources_are_review_not_public(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = settings(tmp_path); write_inputs(cfg, [manifest_item("1-1"), manifest_item("2-2")])
    pipeline = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda i: ready_result(i["document_id"], menu=i["document_id"]))
    report = pipeline.run()
    assert report["documents"]["review"] == 2 and report["records"]["production_ready"] == 0
    assert len(report["conflicts"]) == 1


def test_identical_sources_dedup_preserves_all_provenance(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = settings(tmp_path); write_inputs(cfg, [manifest_item("1-1"), manifest_item("2-2")])
    pipeline = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda i: ready_result(i["document_id"]))
    report = pipeline.run()
    run = cfg.production_root / "runs" / report["run_metadata"]["dataset_version"]
    month = json.loads(next((run / "web_data/menus").rglob("*.json")).read_text(encoding="utf-8"))
    assert month["days"]["2026-01-01"]["breakfast"]["source_document_ids"] == ["1-1", "2-2"]
    assert report["records"]["production_ready"] == 1 and report["records"]["ready_source_records"] == 2


def test_public_schema_tamper_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _, report, run = publish_one(tmp_path, monkeypatch)
    path = next((run / "web_data/menus").rglob("*.json"))
    month = json.loads(path.read_text(encoding="utf-8")); month["days"]["2026-01-01"]["breakfast"]["menu_items"] = [{"name": "other"}]
    atomic_json(path, month)
    with pytest.raises(ValueError, match="projection/index"):
        validate_dataset(run, report, {"I1"})


def test_partial_run_is_not_published(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = settings(tmp_path); write_inputs(cfg, [manifest_item("1-1"), manifest_item("2-2")])
    atomic_json(cfg.production_root / "current.json", {"dataset_version": "previous"})
    pipeline = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda i: ready_result(i["document_id"]))
    pipeline.run(limit=1)
    assert json.loads((cfg.production_root / "current.json").read_text())["dataset_version"] == "previous"


def test_raw_integrity_failure_isolated(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = settings(tmp_path); write_inputs(cfg, [manifest_item("1-1"), manifest_item("2-2")])
    (tmp_path / "raw/1-1.xlsx").write_bytes(b"changed")
    pipeline = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda i: ready_result(i["document_id"]))
    report = pipeline.run()
    assert report["documents"]["ready"] == 1 and report["documents"]["failed"] == 1
    assert report["failure_codes"]["RAW_INTEGRITY_FAILED"] == 1


def test_download_failure_does_not_stop_other_documents(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = settings(tmp_path); items = [manifest_item("1-1"), manifest_item("2-2")]
    items[0].update(download_status="failed", local_path=None, error="timeout")
    write_inputs(cfg, [items[1]])
    atomic_jsonl(cfg.production_manifest_path, items)
    pipeline = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda i: ready_result(i["document_id"]))
    report = pipeline.run()
    assert report["documents"]["ready"] == 1 and report["documents"]["failed"] == 1
    assert report["documents"]["parsed"] == 1 and report["download_failures"] == ["1-1"]


def test_unknown_nonempty_institution_is_review(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = settings(tmp_path); write_inputs(cfg, [manifest_item(institution_id="not-canonical")])
    pipeline = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda i: ready_result(institution_id="not-canonical"))
    report = pipeline.run()
    assert report["documents"]["review"] == 1 and report["web_data"]["institution_month_files"] == 0


def test_recover_crash_after_promoting_run_before_pointer(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = settings(tmp_path); write_inputs(cfg, [manifest_item()])
    pipeline = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda i: ready_result())
    monkeypatch.setattr(pipeline, "_publish_pointer", lambda *args: (_ for _ in ()).throw(OSError("crash")))
    with pytest.raises(OSError):
        pipeline.run()
    assert not (cfg.production_root / "current.json").exists()
    recovery = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(recovery.parser, "parse_sample", lambda i: pytest.fail("promoted run must not be reparsed"))
    report = recovery.run()
    assert json.loads((cfg.production_root / "current.json").read_text())["dataset_version"] == report["run_metadata"]["dataset_version"]


@pytest.mark.parametrize("header,filename_year,expected_warning", [("식단표", None, True), ("2026년 1월 식단표", None, False), ("식단표", 2026, False)])
def test_historical_publication_fallback_requires_confirmation(tmp_path: Path, header: str,
                                                              filename_year: int | None, expected_warning: bool) -> None:
    workbook = WorkbookGrid("unused", "xlsx", "xlsx", False, [
        SheetGrid("Sheet1", 1, 1, 1, 1, cells=[CellData("A1", 1, 1, header, header)])
    ])
    sample = manifest_item()
    sample.update(metadata_year_inferred=True, period_metadata={"filename_year": filename_year, "filename_month": 1})
    _, _, evidence, warnings = ExcelMealParser(tmp_path)._period(sample, workbook)
    assert any(w["code"] == "YEAR_INFERRED_FROM_PUBLICATION" for w in warnings) == expected_warning
    assert evidence


def test_production_runner_uses_real_golden_core(tmp_path: Path) -> None:
    project_root = Path(__file__).resolve().parents[1]
    samples = json.loads((project_root / "tests/fixtures/excel_samples/manifest.json").read_text(encoding="utf-8"))
    golden = json.loads((FIXTURES / "golden_excel/expected.json").read_text(encoding="utf-8"))["documents"]
    golden_ids = {f"{d['post_id']}-{d['attachment_id']}" for d in golden}
    items = []
    for sample in samples:
        document_id = f"{sample['post_id']}-{sample['attachment_id']}"
        if document_id in golden_ids:
            items.append({**sample, "document_id": document_id, "production_status": "pending",
                          "parser_status": "pending", "post_url": f"https://www.corrections.go.kr/policyOpen/corrections/{sample['post_id']}/artclView.do"})
    cfg = replace(settings(tmp_path), project_root=project_root,
                  institutions_path=project_root / "data/institutions/institutions.json")
    atomic_jsonl(cfg.production_manifest_path, items)
    report = ProductionExcelPipeline(cfg).run()
    run = cfg.production_root / "runs" / report["run_metadata"]["dataset_version"]
    for expected in golden:
        result = json.loads((run / "parsed" / f"{expected['post_id']}-{expected['attachment_id']}.json").read_text(encoding="utf-8"))
        actual = {(r["meal_date"], r["meal_type"]): [m["name"] for m in r["menu_items"]] for r in result["records"]}
        for meal in expected["expected"]:
            assert actual[(meal["meal_date"], meal["meal_type"])] == meal["menu_items"]


def test_downloader_failure_checkpoint_and_default_resume(tmp_path: Path) -> None:
    cfg = settings(tmp_path)
    items = [manifest_item("1-1"), manifest_item("2-2")]
    for item in items:
        item.update(download_status="pending", sha256=None, local_path=None)
    atomic_jsonl(cfg.production_manifest_path, items)
    posts = [Post(i["post_id"], "기관", "식단", "2026-01-01", i["post_url"], attachment_count=1, attachments=[
        Attachment(i["attachment_id"], i["filename"], "xlsx", i["download_url"], document_role="inmate")
    ]) for i in items]
    atomic_jsonl(cfg.historical_catalog_path, [p.to_dict() for p in posts])

    class Response:
        headers = {"Content-Type": "application/octet-stream"}
        def iter_content(self, chunk_size: int):
            yield b"source workbook"

    class Client:
        def __init__(self):
            self.calls = []
        def get(self, url: str, stream: bool = False):
            self.calls.append(url)
            if url.endswith("/1"):
                raise TimeoutError("timeout fixture")
            return Response()

    client = Client()
    result = ProductionExcelDownloader(cfg, client).run(limit=1)
    assert result.failed == 1 and result.next_index == 1 and not result.completed
    result = ProductionExcelDownloader(cfg, client).run()
    assert result.failed == 1 and result.success == 1 and result.completed
    assert len(client.calls) == 2
    assert "timeout fixture" in (cfg.production_root / "reports/excel_download_failures.jsonl").read_text(encoding="utf-8")


def test_atomic_json_retries_brief_windows_lock(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import production.io as io
    original_replace = io.os.replace
    calls = []
    def replace_once(source, target):
        calls.append(target)
        if len(calls) == 1:
            raise PermissionError("brief read lock")
        original_replace(source, target)
    monkeypatch.setattr(io.os, "replace", replace_once)
    atomic_json(tmp_path / "value.json", {"complete": True})
    assert len(calls) == 2 and json.loads((tmp_path / "value.json").read_text())["complete"]


def test_bad_public_source_isolated_from_healthy_document(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = settings(tmp_path); items = [manifest_item("1-1"), manifest_item("2-2")]
    items[0]["post_url"] = "data/raw/private.xlsx"
    write_inputs(cfg, items)
    pipeline = ProductionExcelPipeline(cfg)
    monkeypatch.setattr(pipeline.parser, "parse_sample", lambda i: ready_result(i["document_id"]))
    report = pipeline.run()
    assert report["documents"]["review"] == 1 and report["documents"]["ready"] == 1
    assert report["validation_issues"]["PRODUCTION_SCHEMA_INVALID"] == 1


def test_wide_nonmeal_sheet_with_headers_is_not_forced_pass(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from openpyxl.utils import get_column_letter
    import parsers.excel.pipeline as core
    cells = [CellData(f"{get_column_letter(c)}{r}", r, c, "cost", "cost")
             for r in range(3, 103) for c in range(1, 61)]
    cells.extend(CellData(f"{get_column_letter(c)}2", 2, c, label, label)
                 for c, label in [(55, "아침"), (57, "점심"), (59, "저녁")])
    workbook = WorkbookGrid("unused", "xlsx", "xlsx", False, [SheetGrid("costs", 2, 102, 1, 60, cells=cells)])
    monkeypatch.setattr(core, "read_workbook", lambda *args: workbook)
    result = ExcelMealParser(tmp_path).parse_sample(manifest_item())
    assert result["status"] == "FAIL" and result["records_generated"] == 0
    assert result["failures"][0]["code"] == "UNSUPPORTED_LAYOUT"

from pathlib import Path

from openpyxl import Workbook

from audit.structure_survey import survey_file
from audit.template_families import build_template_families


def test_xlsx_structure_survey_reports_formula_and_actual_range(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "식단"
    sheet["A1"] = "날짜"
    sheet["B1"] = "조식"
    sheet["A2"] = "2026년 1월"
    sheet["B2"] = "쌀밥"
    sheet["C2"] = "=1+1"
    sheet.merge_cells("A3:C3")
    workbook.save(path)

    result = survey_file(path)

    assert result["survey_status"] == "ok"
    first = result["sheets"][0]  # type: ignore[index]
    assert first["formula_cells"] == 1
    assert first["nonempty_used_range"]["max_column"] == 3
    assert first["merged_cell_ranges"] == 1


def test_template_fingerprint_groups_equivalent_pdf_samples() -> None:
    records = [
        {
            "survey_status": "ok",
            "extension": "pdf",
            "filename": f"sample-{index}.pdf",
            "institution_id": f"I{index}",
            "page_count": 1,
            "text_character_count": 900,
            "image_count": 0,
            "possible_image_only": False,
            "pages": [{"width_points": 595.28, "height_points": 841.89, "text_block_count": 50}],
        }
        for index in (1, 2)
    ]

    families = build_template_families(records)

    assert len(families["pdf"]) == 1
    assert families["pdf"][0]["sample_count"] == 2


def test_xls_extension_with_xlsx_container_uses_signature_fallback(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "source.xlsx"
    mislabeled_path = tmp_path / "mislabeled.xls"
    workbook = Workbook()
    workbook.active["A1"] = "수용자 식단표"
    workbook.save(xlsx_path)
    mislabeled_path.write_bytes(xlsx_path.read_bytes())

    result = survey_file(mislabeled_path)

    assert result["survey_status"] == "ok"
    assert result["detected_format"] == "xlsx"
    assert result["format_extension_mismatch"] is True
